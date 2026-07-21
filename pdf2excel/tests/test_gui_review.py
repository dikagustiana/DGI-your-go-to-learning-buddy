"""Offscreen tests for the QA review view.

These build synthetic PageResults — no OCR engine, no PDF — and verify
that user actions in the widget mutate the model the way the export
path expects. Skipped automatically when PySide6 isn't installed.
"""

import os
from decimal import Decimal

import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication  # noqa: E402

from app.pipeline.models import Cell, PageResult  # noqa: E402


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture()
def view(qapp):
    from app.gui.review_view import ReviewView
    v = ReviewView()
    yield v
    v.deleteLater()


def make_result():
    return PageResult(
        page_number=1, rotation_applied=90, rotation_source="auto",
        grid=[
            [Cell.from_text("SALDO", 0.95), Cell.from_text("92.240", 0.55)],
            [Cell.from_text("BIAYA", 0.90), Cell.from_text("4.488,00", 0.92)],
        ])


def test_populate_shows_grid_and_rotation(view):
    result = make_result()
    view.set_page(None, 1, result)   # no pdf path: table only, no image
    assert view.table.rowCount() == 2
    assert view.table.columnCount() == 2
    assert view.table.item(0, 1).text() == "92.240"
    assert view.rotation_combo.currentText() == "90"
    assert not view.stale_banner.isVisible()


def test_edit_writes_back_and_reparses(view):
    result = make_result()
    view.set_page(None, 1, result)
    changed = []
    view.result_changed.connect(changed.append)

    # The observed real-world failure: OCR read 92.240, truth is 192.240.
    view.table.item(0, 1).setText("192.240")

    cell = result.grid[0][1]
    assert cell.raw == "192.240"
    assert cell.value == Decimal("192240")   # re-parsed with id-ID rules
    assert cell.edited is True
    assert changed, "result_changed must fire so the session can persist"


def test_edit_to_non_number_clears_value(view):
    result = make_result()
    view.set_page(None, 1, result)
    view.table.item(0, 1).setText("tidak terbaca")
    assert result.grid[0][1].value is None
    assert result.grid[0][1].edited is True


def test_edit_into_empty_cell_creates_cell(view):
    result = make_result()
    result.grid[0][1] = None
    view.set_page(None, 1, result)
    view.table.item(0, 1).setText("500,25")
    cell = result.grid[0][1]
    assert cell is not None
    assert cell.value == Decimal("500.25")
    assert cell.edited and cell.confidence == 1.0


def test_label_and_reviewed_write_back(view):
    result = make_result()
    view.set_page(None, 1, result)
    view.label_combo.setCurrentText("rekening koran")
    assert result.doc_type_label == "rekening koran"
    view.reviewed_check.setChecked(True)
    assert result.reviewed is True


def test_rotation_change_marks_grid_stale_and_requests_reocr(view, qapp):
    result = make_result()
    # A pdf path must be bound for the re-OCR button to be enabled; a
    # missing file only degrades the image preview, which is caught.
    view.set_page("nonexistent.pdf", 1, result)
    requests = []
    view.request_reocr.connect(lambda p, r: requests.append((p, r)))

    # isHidden() reflects the widget's own state even while the window
    # itself is not shown (offscreen test).
    view.rotation_combo.setCurrentText("270")
    assert not view.stale_banner.isHidden()

    view.reocr_btn.click()
    assert requests == [(1, 270)]

    # back to the extraction rotation → banner clears
    view.rotation_combo.setCurrentText("90")
    assert view.stale_banner.isHidden()


def test_has_edits(view):
    result = make_result()
    view.set_page(None, 1, result)
    assert not view.has_edits()
    view.table.item(1, 0).setText("BIAYA ADM")
    assert view.has_edits()


def test_export_colors_edited_over_lowconf(tmp_path):
    from openpyxl import load_workbook
    from app.config import PipelineConfig
    from app.export.excel import export_workbook

    result = make_result()
    result.grid[0][1].edited = True         # low conf (0.55) but corrected
    out = tmp_path / "o.xlsx"
    export_workbook([result], str(out), PipelineConfig())
    ws = load_workbook(str(out))["p1"]
    fills = {c.coordinate: c.fill.start_color.rgb for row in ws.iter_rows()
             for c in row if c.value is not None}
    # Row 1 holds the raw-column headers; the grid starts at row 2.
    assert fills["B2"] == "00D4EDDA"         # green (edited) wins over amber
