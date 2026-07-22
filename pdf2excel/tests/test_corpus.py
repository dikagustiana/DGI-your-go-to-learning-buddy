"""Adversarial-corpus tests with CELL-LEVEL assertions (Fase 6).

These run REAL OCR over the synthetic adversarial pages and assert on
the resulting values/types — not merely that a workbook exists. Marked
``slow`` (a few seconds per page); run in CI and with
``pytest -m slow``. Skipped when the RapidOCR models aren't available.

Synthetic ≠ finance-grade: clean fonts and controlled noise. The real
Canon-scan validation is an on-prem experiment (see the final report);
this guards the pipeline's handling of the dangerous SHAPES.
"""

import json
import os

import pytest

from app.pipeline.engines.rapidocr_engine import RapidOCREngine

pytestmark = pytest.mark.slow

if not RapidOCREngine.is_available():
    pytest.skip("RapidOCR models not available", allow_module_level=True)


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    out = tmp_path_factory.mktemp("adv")
    import samples.make_adversarial as gen
    gen.main(str(out))
    with open(os.path.join(str(out), "expected.json")) as f:
        manifest = json.load(f)
    return str(out), {m["name"]: m for m in manifest}


def _convert(outdir, name):
    from app.config import PipelineConfig
    from app.pipeline.runner import process_pdf
    pages = process_pdf(os.path.join(outdir, f"{name}.pdf"), PipelineConfig())
    assert pages and not pages[0].error, name
    return pages[0]


def _texts(page):
    return {c.effective_text for row in page.grid for c in row if c}


def _numeric_texts(page):
    return {c.effective_text for row in page.grid for c in row
            if c and c.value is not None}


def test_upright_statement_numbers(corpus):
    outdir, exp = corpus
    page = _convert(outdir, "statement_upright")
    numeric = _numeric_texts(page)
    for token in exp["statement_upright"]["numeric"]:
        assert token in numeric, f"{token} should parse as a number"


def test_rotation_corrected(corpus):
    outdir, _ = corpus
    page = _convert(outdir, "statement_rot90")
    # A 90°-rotated scan must be auto-corrected: its amounts read.
    assert page.rotation_applied in (90, 270)
    assert any(c.value is not None for row in page.grid for c in row if c)


def test_leading_zero_accounts_stay_text(corpus):
    outdir, exp = corpus
    page = _convert(outdir, "identifiers")
    for ident in exp["identifiers"]["must_be_text"]:
        cells = [c for row in page.grid for c in row
                 if c and c.effective_text.replace(" ", "") == ident]
        if cells:  # OCR may split; require that where present it is text
            assert all(c.value is None for c in cells), \
                f"{ident} must not be parsed as a number"


def test_adversarial_values_never_misparsed(corpus):
    outdir, _ = corpus
    page = _convert(outdir, "adversarial_values")
    for row in page.grid:
        for c in row:
            if c is None:
                continue
            t = c.effective_text.replace(" ", "")
            if t == "1,234":
                assert c.value is None            # ambiguous, refused
            if t.startswith("=SUM"):
                assert c.value is None            # formula stays text


def test_formula_injection_neutralized_in_export(corpus, tmp_path):
    from openpyxl import load_workbook
    from app.config import PipelineConfig
    from app.export.excel import export_workbook
    outdir, _ = corpus
    page = _convert(outdir, "adversarial_values")
    out = tmp_path / "adv.xlsx"
    export_workbook([page], str(out), PipelineConfig())
    wb = load_workbook(str(out))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert c.data_type == "s"     # never a live formula


def test_blank_page_survives(corpus):
    outdir, _ = corpus
    page = _convert(outdir, "blank_page")
    assert not page.error          # a near-empty page must not crash
