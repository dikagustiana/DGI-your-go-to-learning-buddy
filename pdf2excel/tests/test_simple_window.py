"""Offscreen tests for the simplified elderly-friendly window.

Real OCR never runs here: results are injected directly (or via the
worker slots) so the tests stay fast. The end-to-end OCR path is
covered by the smoke script and CI.
"""

import os

import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

from app.gui.simple_window import (  # noqa: E402
    SimpleMainWindow, friendly_error, unique_output_path)
from app.pipeline.models import Cell, PageResult  # noqa: E402
from app.session import SessionStore  # noqa: E402


@pytest.fixture(scope="session")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture()
def blank_pdf(tmp_path):
    import fitz
    path = str(tmp_path / "laporan keuangan.pdf")
    doc = fitz.open()
    doc.new_page()
    doc.new_page()
    doc.save(path)
    doc.close()
    return path


def result_for(pno):
    return PageResult(page_number=pno, rotation_applied=0,
                      rotation_source="auto",
                      grid=[[Cell.from_text("192.240", 0.9)]])


def make_window(qapp, tmp_path, monkeypatch,
                answer=QMessageBox.StandardButton.Yes):
    monkeypatch.setattr(QMessageBox, "question",
                        staticmethod(lambda *a, **k: answer))
    monkeypatch.setattr(QMessageBox, "exec",
                        lambda self, *a, **k: QMessageBox.StandardButton.Ok)
    out = tmp_path / "Dokumen"
    out.mkdir(exist_ok=True)
    return SimpleMainWindow(output_dir=str(out))


# ------------------------------------------------------------- helpers

def test_unique_output_path(tmp_path):
    p1 = unique_output_path(str(tmp_path), "laporan")
    assert p1.endswith("laporan.xlsx")
    open(p1, "w").close()
    p2 = unique_output_path(str(tmp_path), "laporan")
    assert p2.endswith("laporan (2).xlsx")
    open(p2, "w").close()
    assert unique_output_path(str(tmp_path), "laporan").endswith("laporan (3).xlsx")
    assert unique_output_path(str(tmp_path), "  ").endswith("hasil.xlsx")


def test_friendly_error_never_leaks_jargon():
    msg, detail = friendly_error(PermissionError("Permission denied: x.xlsx"))
    assert "Tutup dulu" in msg
    msg2, _ = friendly_error(Exception("fitz cannot open broken stream"),
                             context="open")
    assert ".pdf" in msg2
    msg3, detail3 = friendly_error(Exception("Traceback ... ValueError"))
    for jargon in ("OCR", "engine", "DPI", "Traceback"):
        assert jargon not in msg3
    assert "ValueError" in detail3      # tapi detail teknis tetap tersimpan


# ------------------------------------------------------------ alur UI

def test_initial_state(qapp, tmp_path, monkeypatch):
    win = make_window(qapp, tmp_path, monkeypatch)
    try:
        assert win.open_btn.isEnabled()
        assert not win.convert_btn.isEnabled()
        assert win.progress.isHidden()
        assert win.cancel_btn.isHidden()
        assert win.open_excel_btn.isHidden()
        assert win.review_btn.isHidden()
    finally:
        win.close()


def test_load_pdf_shows_page_count(qapp, tmp_path, monkeypatch, blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch)
    try:
        win.load_pdf(blank_pdf)
        assert win.convert_btn.isEnabled()
        assert "2 halaman" in win.file_label.text()
        assert "laporan keuangan.pdf" in win.file_label.text()
    finally:
        win.close()


def test_missing_pages_and_resume_flow(qapp, tmp_path, monkeypatch, blank_pdf):
    with SessionStore.for_pdf(blank_pdf) as s:
        s.save_page(result_for(1))
    win = make_window(qapp, tmp_path, monkeypatch)   # jawab Ya → lanjutkan
    try:
        win.load_pdf(blank_pdf)
        assert set(win._results) == {1}
        assert win._missing_pages() == [2]           # hanya sisa yang diproses
        assert "1 halaman sudah selesai" in win.status_label.text()
    finally:
        win.close()


def test_unreviewed_finish_is_draft(qapp, tmp_path, monkeypatch, blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch,
                      answer=QMessageBox.StandardButton.No)
    try:
        win.load_pdf(blank_pdf)
        win._on_page_done(result_for(1))       # not reviewed
        win._on_page_done(result_for(2))
        win._export_and_finish()

        assert win._out_path.startswith(str(tmp_path / "Dokumen"))
        assert "(DRAF)" in os.path.basename(win._out_path)
        assert "DRAF" in win.done_label.text()
        assert "belum diperiksa" in win.done_label.text().lower()
        assert not win.open_excel_btn.isHidden()
        assert not win.review_btn.isHidden()

        from openpyxl import load_workbook
        wb = load_workbook(win._out_path)
        assert wb["p1"]["A1"].value == 192240        # angka id-ID terparse
        assert "Ringkasan" in wb.sheetnames
        assert "Audit" in wb.sheetnames
        summary = " ".join(str(c.value) for row in wb["Ringkasan"].iter_rows()
                           for c in row if c.value)
        assert "DRAF" in summary
    finally:
        win.close()


def test_reviewed_finish_is_final(qapp, tmp_path, monkeypatch, blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch,
                      answer=QMessageBox.StandardButton.No)
    try:
        win.load_pdf(blank_pdf)
        for pno in (1, 2):
            r = result_for(pno)
            r.set_reviewed(True)              # human signed off
            win._on_page_done(r)
        win._export_and_finish()

        assert "(DRAF)" not in os.path.basename(win._out_path)
        assert win._out_path.endswith("laporan keuangan.xlsx")
        assert "FINAL" in win.done_label.text()
    finally:
        win.close()


def test_failed_page_never_shows_as_success(qapp, tmp_path, monkeypatch,
                                            blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch,
                      answer=QMessageBox.StandardButton.No)
    try:
        win.load_pdf(blank_pdf)
        r = result_for(1)
        r.set_reviewed(True)
        win._on_page_done(r)
        bad = PageResult(page_number=2, rotation_applied=0,
                         rotation_source="error", error="boom")
        win._on_page_done(bad)
        win._export_and_finish()
        # A failed page forces DRAF even though the other page is reviewed.
        assert "(DRAF)" in os.path.basename(win._out_path)
        assert "FINAL" not in win.done_label.text()
        assert "gagal dibaca" in win.done_label.text()
        assert "halaman 2" in win.done_label.text()
    finally:
        win.close()


def test_results_persist_via_page_done(qapp, tmp_path, monkeypatch, blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch,
                      answer=QMessageBox.StandardButton.No)
    try:
        win.load_pdf(blank_pdf)
        win._on_page_done(result_for(1))
    finally:
        win.close()
    with SessionStore.for_pdf(blank_pdf) as s:      # tersimpan ke sesi
        assert 1 in s.load_pages()


def test_no_jargon_on_main_screen(qapp, tmp_path, monkeypatch, blank_pdf):
    win = make_window(qapp, tmp_path, monkeypatch)
    try:
        win.load_pdf(blank_pdf)
        texts = " ".join([
            win.windowTitle(), win.open_btn.text(), win.convert_btn.text(),
            win.file_label.text(), win.status_label.text(),
            win.cancel_btn.text(), win.review_btn.text(),
            win.open_excel_btn.text(), win.open_folder_btn.text(),
        ])
        for jargon in ("OCR", "engine", "Engine", "DPI", "rotation",
                       "bounding", "confidence"):
            assert jargon not in texts, jargon
    finally:
        win.close()
