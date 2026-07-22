"""Regression tests for the finance audit-trail guarantees (Fase 1)."""

from decimal import Decimal

from openpyxl import load_workbook

from app.config import PipelineConfig
from app.export.excel import export_workbook
from app.export.policy import final_eligibility
from app.pipeline.models import Anomaly, Cell, PageResult


def page(pno=1, cells=None, reviewed=False, error="", label=""):
    grid = cells if cells is not None else [[Cell.from_text("100.000", 0.9)]]
    p = PageResult(page_number=pno, rotation_applied=0, rotation_source="auto",
                   grid=grid, error=error, doc_type_label=label, dpi=300,
                   engine="rapidocr", config_hash="abc123")
    p.set_reviewed(reviewed)
    return p


def test_ocr_original_survives_edit():
    cell = Cell.from_text("92.240", 0.5)
    cell.apply_correction("192.240", reviewer="andi")
    assert cell.ocr_original == "92.240"          # never overwritten
    assert cell.corrected_text == "192.240"
    assert cell.effective_text == "192.240"
    assert cell.value == Decimal("192240")
    assert cell.edited and cell.edited_at and cell.reviewer == "andi"


def test_audit_sheet_holds_both_texts(tmp_path):
    cell = Cell.from_text("92.240", 0.5)
    cell.apply_correction("192.240", reviewer="andi")
    out = tmp_path / "o.xlsx"
    export_workbook([page(cells=[[cell]], reviewed=True)], str(out),
                    PipelineConfig(), source_pdf="rk.pdf")
    wb = load_workbook(str(out))
    assert "Audit" in wb.sheetnames
    audit = wb["Audit"]
    headers = [c.value for c in audit[1]]
    row = [c.value for c in audit[2]]
    record = dict(zip(headers, row))
    assert record["OCR asli"] == "92.240"         # immutable original present
    assert record["koreksi manusia"] == "192.240"
    assert record["nilai terparse"] == "192240"
    assert record["pemeriksa"] == "andi"
    assert record["bbox (px @dpi)"] is None or "@300dpi" in str(record["bbox (px @dpi)"])


def test_formula_injection_is_neutralized(tmp_path):
    # OCR text starting with "=" must never execute as an Excel formula.
    evil = Cell.from_text("=1+2", 0.9)
    out = tmp_path / "o.xlsx"
    export_workbook([page(cells=[[evil]], reviewed=True)], str(out),
                    PipelineConfig())
    wb = load_workbook(str(out))
    target = wb["p1"]["A1"]
    assert target.data_type == "s"               # stored as string literal
    assert target.value == "=1+2"


def test_long_identifier_stays_text(tmp_path):
    # Account number: 16 digits, leading zero → must not become a float.
    acct = Cell.from_text("0012345678901234", 0.9)
    assert acct.value is None                      # parser refuses
    out = tmp_path / "o.xlsx"
    export_workbook([page(cells=[[acct]], reviewed=True)], str(out),
                    PipelineConfig())
    val = load_workbook(str(out))["p1"]["A1"].value
    assert val == "0012345678901234"               # exact text, no rounding


def test_merged_by_label_does_not_hide_columns(tmp_path):
    # Two pages, different column counts, same label. The old hidden
    # raw-column block could shadow real data; the new layout must show
    # every visible cell of both pages.
    p1 = page(1, cells=[[Cell.from_text("A", 0.9), Cell.from_text("B", 0.9)]],
              reviewed=True, label="faktur")
    p2 = page(2, cells=[[Cell.from_text("C", 0.9), Cell.from_text("D", 0.9),
                         Cell.from_text("E", 0.9)]], reviewed=True,
              label="faktur")
    out = tmp_path / "o.xlsx"
    cfg = PipelineConfig(output_layout="merged_by_label")
    export_workbook([p1, p2], str(out), cfg)
    ws = load_workbook(str(out))["faktur"]
    seen = {c.value for row in ws.iter_rows() for c in row if c.value}
    for letter in ("A", "B", "C", "D", "E"):
        assert letter in seen
    assert not any(d.hidden for d in ws.column_dimensions.values())


def test_final_blocked_by_unreviewed():
    e = final_eligibility([page(1, reviewed=True), page(2, reviewed=False)])
    assert not e.final_ok
    assert any("belum diperiksa" in r for r in e.reasons)


def test_final_blocked_by_failed_page():
    e = final_eligibility([page(1, reviewed=True), page(2, error="boom")])
    assert not e.final_ok
    assert any("gagal" in r for r in e.reasons)


def test_final_blocked_by_blocker_anomaly():
    p = page(1, reviewed=True)
    p.anomalies = [Anomaly(code="ambiguous-number", message="cek",
                           severity="blocker", row=0, col=0)]
    e = final_eligibility([p])
    assert not e.final_ok


def test_final_ok_when_all_clean():
    e = final_eligibility([page(1, reviewed=True), page(2, reviewed=True)])
    assert e.final_ok and not e.reasons


def test_needs_reextraction_blocks_final():
    p = page(1, reviewed=True)
    p.needs_reextraction = True
    assert not final_eligibility([p]).final_ok
